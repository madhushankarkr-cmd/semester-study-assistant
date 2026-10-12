from pptx import Presentation
import os

notes_folder = r"C:\Users\Hari Shankar Kumar\OneDrive\Desktop\semester-study-assistant\data\Descrete_Mathematics\notes"

for file in os.listdir(notes_folder):
    if file.endswith(".pptx"):
        ppt_path = os.path.join(notes_folder, file)
        try:
            prs = Presentation(ppt_path)
            full_text = []

            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        full_text.append(shape.text)

            text = "\n".join(full_text)
            print(f"✅ {file} → {len(prs.slides)} slides | Text length: {len(text)} characters")
        except Exception as e:
            print(f"❌ {file} → Error: {e}")