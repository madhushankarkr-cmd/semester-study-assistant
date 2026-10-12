import os
from pypdf import PdfReader
from pptx import Presentation

def load_pdf(file_path):
    """PDF se text nikalta hai"""
    try:
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        return text.strip()
    except Exception as e:
        print(f"❌ Error reading PDF {file_path}: {e}")
        return ""


def load_pptx(file_path):
    """PPTX se text nikalta hai"""
    try:
        prs = Presentation(file_path)
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text.strip()
    except Exception as e:
        print(f"❌ Error reading PPTX {file_path}: {e}")
        return ""


def load_documents(data_folder):
    """
    Saari PDF aur PPTX files load karta hai.
    Returns: list of dictionaries
    """
    documents = []

    for root, dirs, files in os.walk(data_folder):
        for file in files:
            file_path = os.path.join(root, file)
            file_lower = file.lower()

            # Subject name nikalna (folder se)
            relative_path = os.path.relpath(root, data_folder)
            subject = relative_path.split(os.sep)[0] if relative_path != "." else "Unknown"

            text = ""
            file_type = ""

            if file_lower.endswith(".pdf"):
                text = load_pdf(file_path)
                file_type = "pdf"
            elif file_lower.endswith(".pptx"):
                text = load_pptx(file_path)
                file_type = "pptx"
            else:
                continue  # dusre file types skip

            # Agar text nahi mila (scanned PDF) toh skip
            if not text or len(text) < 50:
                print(f"⚠️  Skipped (no text): {file}")
                continue

            documents.append({
                "subject": subject,
                "filename": file,
                "file_type": file_type,
                "text": text,
                "char_count": len(text)
            })

            print(f"✅ Loaded: {subject} → {file} ({len(text)} chars)")

    return documents


# ------------------- TEST -------------------
if __name__ == "__main__":
    data_path = r"C:\Users\Hari Shankar Kumar\OneDrive\Desktop\semester-study-assistant\data"
    
    docs = load_documents(data_path)
    
    print("\n" + "="*50)
    print(f"Total Documents Loaded: {len(docs)}")
    print("="*50)
    
    for doc in docs:
        print(f"{doc['subject']} | {doc['filename']} | {doc['char_count']} chars")